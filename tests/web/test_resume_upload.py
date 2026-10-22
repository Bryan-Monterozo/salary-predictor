from __future__ import annotations

import io
import logging
import re
import subprocess
import zipfile
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter
from pypdf.generic import (
    ArrayObject,
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
    NullObject,
    NumberObject,
    TextStringObject,
)
from werkzeug.datastructures import FileStorage

from salary_predictor.web import create_app, resume_extract
from salary_predictor.web.profile_flow import get_catalog
from salary_predictor.web.resume_extract import (
    MAX_UPLOAD_BYTES,
    ResumeError,
    extract_text,
    process_upload,
)
from salary_predictor.web.resume_suggestions import suggest_profile


@pytest.fixture
def client():
    return create_app({"TESTING": True, "SECRET_KEY": "resume-fixture-secret"}).test_client()


def pdf_file(lines=(), *, blank=False, encrypted=False, pages=1):
    writer = PdfWriter()
    for index in range(pages):
        page = writer.add_blank_page(width=612, height=792)
        if not blank:
            stream = DecodedStreamObject()
            stream.set_data(
                ("BT /F1 12 Tf 40 750 Td " + " ".join(
                    f"({line}) Tj 0 -16 Td" for line in lines if index == 0
                ) + " ET").encode("ascii")
            )
            font = DictionaryObject({
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            })
            page[NameObject("/Resources")] = DictionaryObject({
                NameObject("/Font"): DictionaryObject({
                    NameObject("/F1"): writer._add_object(font),
                }),
            })
            page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt("private-password")
    target = io.BytesIO()
    writer.write(target)
    return target.getvalue()


def docx_file(lines=(), *, paragraphs=0):
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    for _ in range(paragraphs):
        document.add_paragraph("x")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def csrf(page):
    match = re.search(r'name="_csrf_token" value="([^"]+)"', page.get_data(as_text=True))
    assert match
    return match.group(1)


def upload(client, content, name="career.docx"):
    return client.post(
        "/upload",
        data={
            "_csrf_token": csrf(client.get("/upload")),
            "resume": (io.BytesIO(content), name),
        },
        content_type="multipart/form-data",
    )


@pytest.mark.parametrize("extension", ["pdf", "docx"])
def test_resume_upload_to_same_editable_form_then_review_and_result(client, extension):
    lines = [
        "Role: Software Developer",
        "Occupation: SOFTWARE DEVELOPERS",
        "Region: Metro Manila",
        "Experience: 4 years",
        "Education: Bachelor's degree",
        "Skills: Python, SQL",
    ]
    data = pdf_file(lines) if extension == "pdf" else docx_file(lines)
    page = upload(client, data, "candidate." + extension)
    assert page.status_code == 200
    html = page.get_data(as_text=True)
    assert "Résumé suggestions" in html
    assert 'value="Software Developer"' in html
    assert 'name="occupation_psoc_code"' in html
    assert 'name="region_psgc_code"' in html
    assert 'value="4"' in html
    assert "Python, SQL" in html
    assert "Confirm" in html
    token = csrf(page)
    form = {
        "_csrf_token": token,
        "job_title": "Software Developer",
        "occupation_psoc_code": "2512",
        "region_psgc_code": "1300000000",
        "years_experience": "6",
        "education_level": "bachelor",
        "skills": "SQL",
        "employment_context": "",
    }
    reviewed = client.post("/review", data=form)
    assert reviewed.status_code == 200
    assert b"6 years" in reviewed.data
    assert b"SQL" in reviewed.data
    assert b"Python" not in reviewed.data
    result = client.post("/results", data={**form, "_csrf_token": csrf(reviewed)})
    assert result.status_code == 200
    assert b"NOT YET AVAILABLE" in result.data
    with client.session_transaction() as session:
        assert set(session) == {"_csrf_token"}


def test_verified_plain_extraction_from_pdf_and_docx():
    for extension, data in (
        ("pdf", pdf_file(["Role: Engineer", "Experience: 5 years"])),
        ("docx", docx_file(["Role: Engineer", "Experience: 5 years"])),
    ):
        extracted = extract_text(data, "resume." + extension)
        assert "Engineer" in extracted.text
        assert "Experience: 5 years" in extracted.text
        assert extracted.document_type == extension


@pytest.mark.parametrize("content,name,phrase", [
    (b"", "blank.pdf", "Choose a PDF"),
    (b"not a pdf", "false.pdf", "cannot be read"),
    (b"PK\x03\x04garbage", "broken.docx", "cannot be read"),
    (b"%PDF-1.7 rubbish", "broken.pdf", "cannot be read"),
    (b"%PDF-1.7 fake", "malware.exe", "Only PDF and DOCX"),
    (b"PK\x03\x04fake", "fake.pdf", "cannot be read"),
])
def test_missing_corrupt_spoofed_extension_or_signature(client, content, name, phrase):
    result = upload(client, content, name)
    assert result.status_code == 422
    assert phrase in result.get_data(as_text=True)


def test_encrypted_pdf_rejected(client):
    result = upload(client, pdf_file(["Role: Engineer"], encrypted=True), "locked.pdf")
    assert result.status_code == 422
    assert b"Password-protected" in result.data


def test_textless_pdf_explains_no_ocr(client):
    result = upload(client, pdf_file(blank=True), "scanned.pdf")
    assert result.status_code == 422
    assert b"OCR is not supported" in result.data


def test_pdf_page_limit_rejected(client):
    result = upload(client, pdf_file(["Role: Engineer"], pages=9), "large.pdf")
    assert result.status_code == 422
    assert b"no more than 8 pages" in result.data


def test_docx_paragraph_budget_rejected(client):
    result = upload(client, docx_file(paragraphs=1201), "long.docx")
    assert result.status_code == 422
    assert b"1,200 paragraphs" in result.data


def test_docx_expansion_and_traversal_are_blocked(client):
    content = io.BytesIO()
    with zipfile.ZipFile(content, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("_rels/.rels", "<Relationships/>")
        archive.writestr("word/document.xml", "X" * (5 * 1024 * 1024))
        archive.writestr("../escape.txt", "unsafe")
    result = upload(client, content.getvalue())
    assert result.status_code == 422
    assert b"unsafe" in result.data.lower()


def test_docx_embedded_content_rejected(client):
    original = docx_file(["Role: Engineer"])
    content = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original)) as source:
        with zipfile.ZipFile(content, "w", compression=zipfile.ZIP_DEFLATED) as output:
            for item in source.infolist():
                output.writestr(item.filename, source.read(item.filename))
            output.writestr("word/embeddings/invoice.exe", "not allowed")
    result = upload(client, content.getvalue())
    assert result.status_code == 422
    assert b"unsafe" in result.data.lower()


def test_http_upload_limit_returns_413(client):
    result = upload(client, b"x" * (MAX_UPLOAD_BYTES + 20000), "too-large.pdf")
    assert result.status_code == 413
    assert b"2 MB" in result.data


def test_csrf_and_missing_document(client):
    missing_token = client.post("/upload", data={},
                                content_type="multipart/form-data")
    assert missing_token.status_code == 400
    no_file = client.post("/upload", data={"_csrf_token": csrf(client.get("/upload"))},
                          content_type="multipart/form-data")
    assert no_file.status_code == 422


def test_only_explicit_resume_fields_are_suggested(client):
    catalog = get_catalog(str(Path(__file__).resolve().parents[2]))
    draft, notes = suggest_profile(
        "Jane Candidate\njane@example.com\n2021 - 2025\n"
        "Occupation: software developer\nLocation: NCR\n"
        "Experience: 2021-2025\nSkills: SQL, Python", catalog
    )
    assert not draft["job_title"]
    assert not draft["occupation_psoc_code"]
    assert draft["region_psgc_code"] == "1300000000"
    assert not draft["years_experience"]
    assert draft["skills"] == "SQL, Python"
    assert any("ambiguous" in note.lower() or "unambiguous" in note.lower()
               for note in notes)
    assert "jane@example.com" not in str(draft) + str(notes)


def test_explicit_months_and_uncertain_education_do_not_fabricate(client):
    catalog = get_catalog(str(Path(__file__).resolve().parents[2]))
    draft, notes = suggest_profile(
        "Role: Researcher\nExperience: 6 months\nEducation: college",
        catalog,
    )
    assert draft["job_title"] == "Researcher"
    assert draft["years_experience"] == "0.50"
    assert draft["education_level"] == ""
    assert any("not uniquely" in note for note in notes)


def test_suspicious_resume_text_is_not_echoed_in_errors_or_logs(client, caplog):
    private = "really.private@example.com"
    caplog.set_level(logging.INFO)
    content = docx_file(["Role: Engineer", "Skills: SQL", private])
    result = upload(client, content)
    assert result.status_code == 200
    assert private not in result.get_data(as_text=True)
    assert private not in caplog.text
    assert client.get("/profile").status_code == 200


def test_worker_timeout_maps_to_safe_error(monkeypatch):
    def timed_out(*args, **kwargs):
        raise subprocess.TimeoutExpired(args[0], 5)

    monkeypatch.setattr(resume_extract.subprocess, "run", timed_out)
    with pytest.raises(ResumeError, match="too long"):
        process_upload(FileStorage(stream=io.BytesIO(pdf_file(["Role: Doctor"])),
                                   filename="safe.pdf"))


def test_input_name_not_used_as_worker_path(client):
    result = upload(client, docx_file(["Role: Engineer"]), "../../any.docx")
    assert result.status_code == 200
    assert b"Engineer" in result.data


def test_upload_leaves_no_resume_file_in_project_directory(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    client = create_app({"TESTING": True, "SECRET_KEY": "storage-check"}).test_client()
    result = upload(client, docx_file(["Role: Analyst", "Skills: SQL"]))
    assert result.status_code == 200
    assert list(tmp_path.rglob("*")) == []


def test_resume_suggestion_escapes_embedded_html(client):
    result = upload(client, docx_file(["Role: <img src=x onerror=alert(1)>"]))
    assert result.status_code == 200
    html = result.get_data(as_text=True)
    assert "<img src=x" not in html
    assert "&lt;img" in html


def test_date_range_is_shown_as_evidence_not_experience_value():
    catalog = get_catalog(str(Path(__file__).resolve().parents[2]))
    draft, notes = suggest_profile("Role: Analyst\n2021 - 2025", catalog)
    assert draft["years_experience"] == ""
    assert any("2021 - 2025" in note for note in notes)


def spaced_glyphs(line: str) -> str:
    return "  ".join(" ".join(word) for word in line.split(" "))


def test_unlabeled_sectioned_resume_from_pdf_populates_fields(client):
    lines = [
        "EDUCATIONAL BACKGROUND",
        "College (Bachelor of",
        "Science in Information",
        "Technology)",
        "2019 - 2026",
        "SKILLS",
        "Basic Computer",
        "Troubleshooting",
        "Basic Network",
        "Troubleshooting",
        "Good Adaptability Skills",
        "Good Communication Skills",
        "WORK EXPERIENCE",
        "2021-2022",
        "EXAMPLE NETWORK SERVICES",
        "Installer of Internet Cable",
        "October 2024 - January 2025 (Part time)",
        "SAMPLE HOTEL",
        "Banquet Waiter",
        "January 12 2026 - April 28 2026 (OJT)",
        "EXAMPLE SYSTEMS GROUP",
        "IT Support Specialist",
    ]
    response = upload(client, pdf_file([spaced_glyphs(line) for line in lines]), "cv.pdf")
    assert response.status_code == 200
    html = response.get_data(as_text=True)
    assert 'value="IT Support Specialist"' in html
    assert "Basic Computer Troubleshooting, Basic Network Troubleshooting" in html
    assert "Good Adaptability Skills, Good Communication Skills" in html
    assert "actual total years" in html
    assert "marked as OJT" in html
    assert "not be your current job" in html
    assert 'value="1300000000" selected' not in html
    assert '"years_experience" value=""' not in html
    assert '"education_level" value="bachelor"' not in html
    assert "EXAMPLE SYSTEMS GROUP" not in html


def test_unlabeled_sectioned_docx_resume_is_supported(client):
    lines = [
        "EDUCATIONAL BACKGROUND",
        "Bachelor of Science in Information Technology (in progress)",
        "SKILLS",
        "Basic Network",
        "Troubleshooting",
        "Good Communication Skills",
        "WORK EXPERIENCE",
        "2024-2025",
        "A SAMPLE COMPANY",
        "IT Support Specialist",
    ]
    response = upload(client, docx_file(lines), "cv.docx")
    assert response.status_code == 200
    html = response.get_data(as_text=True)
    assert 'value="IT Support Specialist"' in html
    assert "Basic Network Troubleshooting" in html
    assert "Good Communication Skills" in html
    assert "highest completed level" in html


def test_resume_contact_address_is_not_job_region_suggestion(client):
    lines = [
        "CONTACT", "A sample residence in Quezon City", "SKILLS", "Troubleshooting",
        "WORK EXPERIENCE", "2024-2025", "SAMPLE SERVICES", "Support Technician",
    ]
    response = upload(client, docx_file(lines), "cv.docx")
    assert response.status_code == 200
    html = response.get_data(as_text=True)
    assert "contact-city information" in html
    assert 'value="1300000000" selected' not in html


def test_pdf_style_letter_spacing_preserves_degree_notes_without_claiming_graduation():
    raw_lines = [
        "EDUCATIONAL BACKGROUND",
        "College (Bachelor of",
        "Science in Information",
        "Technology)",
        "2019 - 2026",
        "SKILLS",
        "Basic Computer",
        "Troubleshooting",
        "Basic Network",
        "Troubleshooting",
        "Good Communication Skills",
        "WORK EXPERIENCE",
        "2021-2022",
        "EXAMPLE SERVICES GROUP",
        "Installer of Internet Cable",
        "January 2026 - April 2026 (OJT)",
        "SAMPLE COMPANY",
        "IT Support Specialist",
    ]
    text = "\n".join(spaced_glyphs(line) for line in raw_lines)
    catalog = get_catalog(str(Path(__file__).resolve().parents[2]))
    draft, notes = suggest_profile(text, catalog)
    assert draft["job_title"] == "IT Support Specialist"
    assert draft["skills"] == (
        "Basic Computer Troubleshooting, Basic Network Troubleshooting, "
        "Good Communication Skills"
    )
    assert draft["education_level"] == ""
    assert draft["years_experience"] == ""
    assert any("bachelor's program" in note for note in notes)
    assert any("part-time" in note for note in notes)


def pdf_with_open_action(kind: str) -> bytes:
    writer = PdfWriter()
    writer.append_pages_from_reader(PdfReader(io.BytesIO(pdf_file(["Role: Data Analyst"]))))
    root = writer._root_object
    if kind == "internal":
        root[NameObject("/OpenAction")] = ArrayObject([
            writer.pages[0].indirect_reference,
            NameObject("/XYZ"), NullObject(), NullObject(), NumberObject(0),
        ])
        writer.add_attachment("Content Credentials", b"sample harmless metadata")
    elif kind == "script":
        root[NameObject("/OpenAction")] = DictionaryObject({
            NameObject("/S"): NameObject("/JavaScript"),
            NameObject("/JS"): TextStringObject("app.alert('No')"),
        })
    elif kind == "additional":
        root[NameObject("/AA")] = DictionaryObject({
            NameObject("/O"): DictionaryObject({
                NameObject("/S"): NameObject("/JavaScript"),
                NameObject("/JS"): TextStringObject("app.alert('No')"),
            }),
        })
    else:
        raise ValueError("unknown kind")
    data = io.BytesIO()
    writer.write(data)
    return data.getvalue()


def test_pdf_internal_open_destination_with_content_credentials_is_accepted(client):
    payload = pdf_with_open_action("internal")
    result = extract_text(payload, "international.pdf")
    assert "Role: Data Analyst" in result.text
    page = upload(client, payload, "international.pdf")
    assert page.status_code == 200
    assert 'value="Data Analyst"' in page.get_data(as_text=True)


@pytest.mark.parametrize("kind", ["script", "additional"])
def test_pdf_executable_open_actions_still_rejected(client, kind):
    page = upload(client, pdf_with_open_action(kind), "unsafe.pdf")
    assert page.status_code == 422
    assert b"cannot be read" in page.data


@pytest.mark.parametrize("extension", ["pdf", "docx"])
def test_realistic_engineering_cv_headings_and_inline_work_dates(client, extension):
    lines = [
        "PROFESSIONAL SUMMARY",
        "Software engineer working with web and mobile projects.",
        "TECHNICAL SKILLS",
        "Languages & frameworks: Python, JavaScript, HTML",
        "Backend & databases: SQL Server, REST APIs, SQL",
        "Tools & practices: Git, QA, release verification,",
        "DevOps support",
        "PROFESSIONAL EXPERIENCE",
        "SAMPLE SERVICES | Philippines",
        "Software Developer II | Aug 2026-Present",
        "Software Developer I | Mar 2025-Aug 2026",
        "Junior Developer | Aug 2024-Mar 2025",
        "EDUCATION",
        "SAMPLE UNIVERSITY | 2020-2024",
        "Bachelor of Science in Computer Engineering",
    ]
    payload = pdf_file(lines) if extension == "pdf" else docx_file(lines)
    response = upload(client, payload, "sample." + extension)
    assert response.status_code == 200
    html = response.get_data(as_text=True)
    assert 'value="Software Developer II"' in html
    assert "marked as ongoing" in html
    assert "Python, JavaScript, HTML, SQL Server, REST APIs, SQL" in html
    assert "Git, QA, release verification, DevOps support" in html
    assert "bachelor&#39;s program" in html or "bachelor's program" in html
    assert "actual total years" in html
    assert 'value="1300000000" selected' not in html
